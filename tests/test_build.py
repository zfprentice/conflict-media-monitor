"""End to end: one news refresh against a fake Google News, then an offline rebuild of every output."""
import json
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

import build_data
from monitor import archive, events

ROOT = Path(__file__).parent.parent
FULL = json.loads((ROOT / "config" / "countries.json").read_text())
NOW = datetime(2026, 10, 21, 5, 0, tzinfo=timezone.utc)


def small_config(tmp_path):
    cfg = {**FULL, "countries": [c for c in FULL["countries"] if c["iso3"] in {"SDN", "LBN"}]}
    cfg["watch_terms"] = [{**FULL["watch_terms"][0], "backfill_from": "2026-09-01",
                           "queries": FULL["watch_terms"][0]["queries"][:2]}]
    p = tmp_path / "countries.json"
    p.write_text(json.dumps(cfg))
    return cfg, p


def fake_fetch(calls):
    def fetch(query, deadline=None, window="2d", edition="GB:en", after=None, before=None, **_):
        calls.append((query, window, edition, after))
        when = (datetime.fromisoformat(after) + timedelta(days=1, hours=9)).replace(tzinfo=timezone.utc) if after \
            else NOW - timedelta(hours=3)
        stamp = when.isoformat()
        if "Institut" in query:
            return [{"title": f"Mediation in Yemen: {edition} {after}", "url": f"https://n/{edition}{after}",
                     "source": "Le Monde" if "FR" in edition else "National Post", "published": stamp}]
        if "Sudan" in query:
            return [
                {"title": "Drone strike on Abyei market kills 150, Sudan officials say", "url": "https://n/1", "source": "Reuters", "published": stamp},
                {"title": "Sudanese drone attack on Abyei market kills more than 150", "url": "https://n/2", "source": "BBC", "published": stamp},
                {"title": "Korea expands farm support for South Sudan", "url": "https://n/3", "source": "Korea Times", "published": stamp},
                {"title": "Hot, humid day ahead in Khartoum with drizzle", "url": "https://n/4", "source": "X", "published": stamp},
            ]
        return [{"title": "Israel strikes south Lebanon", "url": "https://n/5", "source": "AFP", "published": stamp}]
    return fetch


def test_refresh_then_offline_build(tmp_path, monkeypatch):
    cfg, cfg_path = small_config(tmp_path)
    data = tmp_path / "data"
    data.mkdir()
    p = build_data.paths(data)
    calls = []
    monkeypatch.setattr(build_data.gnews, "fetch", fake_fetch(calls))
    monkeypatch.setattr(build_data.time, "sleep", lambda s: None)

    news, mentions = {}, {"eip": []}
    stats = build_data.refresh_news(cfg, p, NOW, news, mentions, build_data.NewsClient(60))
    assert stats["news_failed"] == 0

    # headlines: off-topic and noise dropped, the two Abyei reports grouped into one story
    sdn = news["SDN"]
    assert [h["url"] for h in sdn["latest"]] == ["https://n/1", "https://n/2"]
    assert sdn["recent"]["stories"][0]["outlets"] == 2
    # this week's snapshot is open; past weeks were searched by date range and frozen
    assert not archive.load_week(p["weeks"], "2026-W43")["final"]
    w42 = archive.load_week(p["weeks"], "2026-W42")
    assert w42["final"] and w42["countries"]["SDN"]["stories"][0]["outlets"] == 2
    assert any(a == "2026-10-12" for _, _, _, a in calls)
    # watch term: last 7 days plus month-by-month backfill in both editions, de-duplicated
    titles = {m["title"] for m in mentions["eip"]}
    assert "Mediation in Yemen: FR:fr 2026-09-01" in titles and "Mediation in Yemen: GB:en None" in titles
    done = p["backfill"].read_text().split()
    assert "eip|English|2026-09-01" in done and "eip|English|2026-10-01" not in done  # current month is searched again

    # a second refresh adds nothing new and doesn't re-search frozen weeks
    calls.clear()
    build_data.refresh_news(cfg, p, NOW, news, mentions, build_data.NewsClient(60))
    assert not any(a and a < "2026-10-01" for _, _, _, a in calls)
    n_before = len(mentions["eip"])
    assert len({m["title"] for m in mentions["eip"]}) == n_before

    # write the records the way main() does, add some GDELT history, then rebuild offline
    p["news"].write_text(json.dumps(news))
    archive.save_mentions(data / "eip_mentions.csv", mentions["eip"])
    agg = defaultdict(lambda: defaultdict(float))
    for i in range(40):
        d = (date(2026, 9, 11) + timedelta(days=i)).isoformat()
        agg[(d, events.ALL)].update(slots=96, articles=100000, events=5000)
        agg[(d, "SDN")].update(articles=300 + (i * 37 % 11) * 6 + (900 if i == 39 else 0), events=20, tone_w=-900, viol=60)
        agg[(d, "LBN")].update(articles=400, events=20, tone_w=-800, viol=40, peace=50)
    events.save_daily(p["daily"], agg, "")
    p["run"].write_text(json.dumps({"generated_at": NOW.isoformat(timespec="seconds"), "news_failed": 0, "news_total": 4}))

    monkeypatch.setenv("OFFLINE", "1")
    assert build_data.main(data, cfg_path) == 0
    latest = json.loads(p["latest"].read_text())
    sdn = next(c for c in latest["countries"] if c["iso3"] == "SDN")
    assert sdn["attention"]["key"] == "much_higher" and sdn["summary"][0].startswith("Much more attention than usual")
    assert sdn["recent"]["stories"][0]["outlets"] == 2 and sdn["headlines"]
    assert latest["watch"][0]["total"] == n_before and latest["watch"][0]["backfill_total"] == 2
    weekly = json.loads(p["weekly"].read_text())["weeks"]
    assert weekly[0]["week"] == "2026-W43" and weekly[0]["current"]
    w42 = next(w for w in weekly if w["week"] == "2026-W42")
    assert w42["countries"]["SDN"]["collected"] and w42["countries"]["SDN"]["stats"]["days"] == 7
    assert len(json.loads((data / "eip_mentions.json").read_text())) == n_before
    wb = load_workbook(data / "eip_mentions.xlsx")
    assert wb["Mentions"].max_row == n_before + 1 and wb["Mentions"]["B2"].hyperlink
    hist = load_workbook(p["history_xlsx"])
    assert hist["Daily by country"].max_row == 1 + 2 * 40
    assert not p["run"].read_text().count("history_days")  # offline runs never rewrite the run log
